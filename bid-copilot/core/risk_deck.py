"""Génère la présentation de Comité de Risque (.pptx) à partir des données de l'offre.

Séquence : Projet → Chiffres clés → Planning → Risques → Opportunités →
Hypothèses → Synthèse financière → Décisions attendues.
"""
from io import BytesIO

import numpy as np
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
ACCENT = RGBColor(0x00, 0x8C, 0xC8)
RED = RGBColor(0xC0, 0x39, 0x2B)
ORANGE = RGBColor(0xE6, 0x7E, 0x22)
GREEN = RGBColor(0x27, 0xAE, 0x60)
GREY = RGBColor(0x55, 0x5F, 0x6B)
GRAVITE_COLOR = {"Critique": RED, "Élevée": ORANGE, "Moyenne": ACCENT, "Faible": GREEN}


def _keur(v: float) -> str:
    return f"{v / 1000:,.1f} M€".replace(",", " ")


class Deck:
    def __init__(self, titre_projet: str):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.titre_projet = titre_projet
        self.n = 0

    def slide(self, titre: str):
        self.n += 1
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        band = s.shapes.add_shape(1, 0, 0, self.prs.slide_width, Inches(1.0))
        band.fill.solid(); band.fill.fore_color.rgb = NAVY; band.line.fill.background()
        tf = band.text_frame
        tf.text = f"{self.n}. {titre}"
        p = tf.paragraphs[0]; p.font.size = Pt(28); p.font.bold = True; p.font.color.rgb = RGBColor(255, 255, 255)
        foot = s.shapes.add_textbox(Inches(0.5), Inches(7.0), Inches(12), Inches(0.4)).text_frame
        foot.text = f"{self.titre_projet} · Comité de Risque · Document de travail généré automatiquement – à valider"
        foot.paragraphs[0].font.size = Pt(10); foot.paragraphs[0].font.color.rgb = GREY
        return s

    def bullets(self, s, items, left=0.6, top=1.3, width=12, height=5.5, size=18):
        tf = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height)).text_frame
        tf.word_wrap = True
        for i, it in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            is_header = it.isupper()
            p.text = it if is_header else f"• {it}"
            p.font.size = Pt(size); p.space_after = Pt(8); p.font.bold = is_header
            if is_header:
                p.font.color.rgb = ACCENT
        return tf

    def table(self, s, header, rows, top=1.3, col_widths=None, size=12):
        shape = s.shapes.add_table(len(rows) + 1, len(header), Inches(0.5), Inches(top), Inches(12.3), Inches(0.4 * (len(rows) + 1)))
        t = shape.table
        for j, h in enumerate(header):
            c = t.cell(0, j); c.text = h
            c.fill.solid(); c.fill.fore_color.rgb = NAVY
            c.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
            c.text_frame.paragraphs[0].font.size = Pt(size); c.text_frame.paragraphs[0].font.bold = True
        for i, row in enumerate(rows, start=1):
            for j, v in enumerate(row):
                c = t.cell(i, j); c.text = str(v)
                c.text_frame.paragraphs[0].font.size = Pt(size)
        if col_widths:
            for j, w in enumerate(col_widths):
                t.columns[j].width = Inches(w)
        return t

    def kpi(self, s, x, y, label, value, color=ACCENT):
        box = s.shapes.add_shape(1, Inches(x), Inches(y), Inches(2.9), Inches(1.5))
        box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0xF2, 0xF5, 0xF8); box.line.color.rgb = color
        tf = box.text_frame; tf.text = value
        tf.paragraphs[0].font.size = Pt(28); tf.paragraphs[0].font.bold = True; tf.paragraphs[0].font.color.rgb = color
        p = tf.add_paragraph(); p.text = label; p.font.size = Pt(13); p.font.color.rgb = GREY

    def save(self) -> bytes:
        buf = BytesIO(); self.prs.save(buf); return buf.getvalue()


def build_deck(projet: dict, analyse, est, stats: dict, sims: np.ndarray, phases: list[dict],
               similaires, marge_pct: float = 10.0) -> bytes:
    d = Deck(projet["nom"])
    prix_vente = stats["p80"] * (1 + marge_pct / 100)

    # 1. Projet
    s = d.slide("Projet")
    d.bullets(s, [
        f"Client : {projet['client']}",
        f"Objet : {analyse.resume}",
        f"Type : {analyse.type_projet} · Pays : {analyse.pays}",
        f"Remise de l'offre : {projet['date_remise']}",
        f"Projets de référence : " + ", ".join(similaires["id"].tolist()),
    ])

    # 2. Chiffres clés
    s = d.slide("Chiffres clés")
    d.kpi(s, 0.5, 1.5, "Prix de vente proposé", _keur(prix_vente), NAVY)
    d.kpi(s, 3.6, 1.5, "Coût de base (ratios)", _keur(stats["base"]))
    d.kpi(s, 6.7, 1.5, "Coût P80 (Monte Carlo)", _keur(stats["p80"]), ORANGE)
    d.kpi(s, 9.8, 1.5, "Marge visée", f"{marge_pct:.0f} %", GREEN)
    d.bullets(s, [
        f"Tension {analyse.tension_kv} kV · {analyse.nb_travees} travées · {analyse.puissance_mva} MVA · {analyse.km_cable} km de câble",
        f"Délai exigé par le client : {analyse.delai_exige_mois} mois",
        f"{len(analyse.risques)} risques identifiés dont {sum(r.gravite == 'Critique' for r in analyse.risques)} critiques",
    ], top=3.5)

    # 3. Planning
    s = d.slide("Planning")
    fin = max(p["debut"] + p["duree"] for p in phases)
    x0, w = 3.8, 9.0
    scale = w / fin
    for i, ph in enumerate(phases):
        y = 1.4 + i * 0.7
        lab = s.shapes.add_textbox(Inches(0.4), Inches(y), Inches(3.3), Inches(0.5)).text_frame
        lab.text = ph["phase"]; lab.paragraphs[0].font.size = Pt(13)
        bar = s.shapes.add_shape(1, Inches(x0 + ph["debut"] * scale), Inches(y + 0.05), Inches(ph["duree"] * scale), Inches(0.4))
        critical = "transformateur" in ph["phase"].lower()
        bar.fill.solid(); bar.fill.fore_color.rgb = RED if critical else ACCENT; bar.line.fill.background()
        bar.text_frame.text = f"M{ph['debut']}–M{ph['debut'] + ph['duree']}"
        bar.text_frame.paragraphs[0].font.size = Pt(10); bar.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    if analyse.delai_exige_mois:
        xl = x0 + analyse.delai_exige_mois * scale
        line = s.shapes.add_shape(1, Inches(xl), Inches(1.3), Inches(0.04), Inches(0.7 * len(phases)))
        line.fill.solid(); line.fill.fore_color.rgb = RED; line.line.fill.background()
        note = s.shapes.add_textbox(Inches(0.4), Inches(1.4 + 0.7 * len(phases) + 0.2), Inches(12), Inches(0.6)).text_frame
        note.text = (f"Trait rouge : délai client de {analyse.delai_exige_mois} mois. "
                     f"Durée réaliste estimée : {fin} mois → écart de {fin - analyse.delai_exige_mois} mois.")
        note.paragraphs[0].font.size = Pt(14); note.paragraphs[0].font.bold = True; note.paragraphs[0].font.color.rgb = RED

    # 4. Risques
    s = d.slide("Risques")
    ordre = {"Critique": 0, "Élevée": 1, "Moyenne": 2, "Faible": 3}
    risques = sorted(analyse.risques, key=lambda r: ordre[r.gravite])[:7]
    t = d.table(s, ["Gravité", "Risque", "Catégorie", "Mitigation"],
                [[r.gravite, r.titre, r.categorie, r.mitigation] for r in risques],
                col_widths=[1.3, 4.2, 1.6, 5.2], size=11)
    for i, r in enumerate(risques, start=1):
        t.cell(i, 0).text_frame.paragraphs[0].font.color.rgb = GRAVITE_COLOR[r.gravite]
        t.cell(i, 0).text_frame.paragraphs[0].font.bold = True

    # 5. Opportunités
    s = d.slide("Opportunités")
    d.bullets(s, projet.get("opportunites") or [
        "Extension future à 4 transformateurs : position privilégiée pour la tranche 2",
        "Data centers : segment en forte croissance, référence vitrine en Île-de-France",
        "Variante GIS (poste sous enveloppe métallique) : emprise réduite, argument technique",
        "Achats groupés câbles 225 kV avec les autres affaires en cours",
    ])

    # 6. Hypothèses & exclusions
    s = d.slide("Hypothèses, exclusions et données manquantes")
    d.bullets(s, ["HYPOTHÈSES"] + analyse.hypotheses_proposees, left=0.5, width=6.1, size=14)
    d.bullets(s, ["EXCLUSIONS"] + analyse.exclusions_proposees + ["DONNÉES MANQUANTES"] + analyse.donnees_manquantes,
              left=6.8, width=6.1, size=14)

    # 7. Synthèse financière
    s = d.slide("Synthèse financière")
    rows = [[l["libelle"], _keur(l["estimation"])] for _, l in est.iterrows()]
    rows += [["Coût de base", _keur(stats["base"])], ["Provision pour risques (P80 − base)", _keur(stats["provision_risques"])],
             ["Coût P80", _keur(stats["p80"])], [f"Marge ({marge_pct:.0f} %)", _keur(prix_vente - stats["p80"])],
             ["PRIX DE VENTE", _keur(prix_vente)]]
    d.table(s, ["Poste", "Montant"], rows, col_widths=[3.6, 1.8], size=11)
    s.shapes[-1].width = Inches(5.4)
    hist, edges = np.histogram(sims / 1000, bins=20)
    cd = CategoryChartData()
    cd.categories = [f"{e:.0f}" for e in edges[:-1]]
    cd.add_series("Simulations", hist.tolist())
    ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.3), Inches(1.3), Inches(6.6), Inches(4.6), cd).chart
    ch.has_legend = False
    ch.has_title = True; ch.chart_title.text_frame.text = "Distribution du coût (M€) – 10 000 tirages"
    ch.plots[0].series[0].format.fill.solid(); ch.plots[0].series[0].format.fill.fore_color.rgb = ACCENT
    note = s.shapes.add_textbox(Inches(6.3), Inches(6.0), Inches(6.6), Inches(0.8)).text_frame
    note.text = (f"P50 {_keur(stats['p50'])} · P80 {_keur(stats['p80'])} · P90 {_keur(stats['p90'])} · "
                 f"probabilité de dépasser le coût de base : {stats['proba_depassement_base'] * 100:.0f} %")
    note.paragraphs[0].font.size = Pt(12)

    # 8. Décisions attendues
    s = d.slide("Décisions attendues")
    d.bullets(s, [
        f"Recommandation : {analyse.recommandation_go_nogo}",
        f"Valider un prix de vente de {_keur(prix_vente)} (coût P80 + marge de {marge_pct:.0f} %)",
        "Valider les réserves contractuelles : plafond des pénalités, indexation du cuivre, planning recalé",
        "Autoriser la réservation d'un créneau de fabrication transformateur avant la remise de l'offre",
        "Questions à envoyer au client : " + " / ".join(analyse.questions_au_client[:3]),
    ], size=17)
    return d.save()
